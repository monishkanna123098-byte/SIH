#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
lm_report.py -- the inspection record, and the artifact an officer hands over.

WHAT THIS IS. The place where three different kinds of knowledge meet on one
page without being allowed to look alike:

    EXTRACTED   read off an image by OCR or a vision model. Unverified.
    MEASURED    produced by the metrology engine, carrying an uncertainty.
    DETERMINED  the officer's own call, which overrides both.

A competing entry already ships the two-tier version of this (AI-assisted
finding vs officer determination) and it is good. The middle row is the one
nobody else can populate, because populating it requires a scale reference in
the frame. If a designer ever normalises those three so the table looks
consistent, the product is gone and the page is just another compliance app.

ON THE HASH, AND WHY IT IS NOT A SIGNATURE. `content_hash()` is SHA-256 over a
canonical serialisation of the RECORD, not over the PDF bytes. PDFs embed
creation timestamps, so hashing the file gives a different digest every run and
verifies nothing. Hashing the content means anyone holding the record can
recompute the digest and detect alteration.

It does NOT prove who produced the report, and it does not prevent anyone from
regenerating a different report with a matching hash of its own. It is a
tamper-EVIDENT check against an independently held copy, not a signature, and
the report says exactly that on its face. Calling it "tamper-proof" would be
the kind of claim this project spends its time removing.

ON WHAT IS ABSENT: there is no compliance score and no percentage. See
`lm_declarations.summarise` for why -- CANNOT_DETERMINE cannot be averaged with
PASS, and a single green number is the most efficient way to destroy everything
the three tiers are for.

RUN: python3 lm_report.py        (self-test + writes sample PDF/DOCX)
"""
from __future__ import annotations

import datetime
import hashlib
from dataclasses import dataclass, field
from typing import Optional, Sequence

from lm_declarations import (
    Finding,
    TIER_DETERMINED,
    TIER_EXTRACTED,
    TIER_MEASURED,
    VERDICT_CANNOT_DETERMINE,
    VERDICT_FAIL,
    VERDICT_PASS,
    summarise,
)

VERSION = "1.0-chunk3-2026-09-12"

# Bands the metrology engine can return. Mirrored rather than imported so this
# module does not drag numpy/cv2 in behind it -- a report layer must be able to
# run on a machine that has no imaging stack.
BAND_COMPLIANT = "COMPLIANT"
BAND_DEFICIENT = "DEFICIENT"
BAND_REFER = "REQUIRES_PHYSICAL_VERIFICATION"


@dataclass(frozen=True)
class Measurement:
    """The metrology engine's output, as the report needs it.

    Field-for-field the contract in BUILD-SPEC §5. Everything here already
    exists on `lm_metrology_v7.Verdict` and `lm_capture.manifest_lines()`; this
    dataclass exists so the report layer does not import the imaging stack.

    `height_mm` and `u_mm` are None on a refusal, together, always. A height
    without an uncertainty is the thing this whole project exists to not print.
    """

    band: str
    height_mm: Optional[float] = None
    u_mm: Optional[float] = None
    threshold_mm: Optional[float] = None
    convention: str = ""
    refusal_code: Optional[str] = None
    refusal_detail: str = ""
    threshold_source_tier: str = ""
    manifest: tuple[str, ...] = ()
    roi_box: Optional[tuple[int, int, int, int]] = None

    def __post_init__(self):
        if (self.height_mm is None) != (self.u_mm is None):
            raise ValueError(
                "height_mm and u_mm must both be present or both be None. "
                "A height without an uncertainty is not a measurement.")

    def verdict(self) -> str:
        if self.band == BAND_COMPLIANT:
            return VERDICT_PASS
        if self.band == BAND_DEFICIENT:
            return VERDICT_FAIL
        return VERDICT_CANNOT_DETERMINE

    def stated(self) -> str:
        """The measured value as it may appear in front of anyone.

        Two decimal places, because not every budget term is measured and
        further digits would be arithmetic rather than accuracy -- the same
        rule the metrology engine's own report applies.
        """
        if self.height_mm is None:
            return "not measured"
        return f"{self.height_mm:.2f} mm ± {self.u_mm:.2f} mm (k=2)"

    def band_explanation(self) -> str:
        if self.refusal_code:
            return (f"{self.refusal_code}: {self.refusal_detail}"
                    if self.refusal_detail else self.refusal_code)
        if self.height_mm is None:
            return "No measurement was produced."
        lo, hi = self.height_mm - self.u_mm, self.height_mm + self.u_mm
        tl = self.threshold_mm
        if tl is None:
            return (f"Measured band [{lo:.2f}, {hi:.2f}] mm. No threshold was "
                    f"resolved, so no comparison was made.")
        if self.band == BAND_COMPLIANT:
            return (f"The whole measured band [{lo:.2f}, {hi:.2f}] mm lies at "
                    f"or above the {tl:.2f} mm requirement.")
        if self.band == BAND_DEFICIENT:
            return (f"The whole measured band [{lo:.2f}, {hi:.2f}] mm lies "
                    f"below the {tl:.2f} mm requirement.")
        return (f"The measured band [{lo:.2f}, {hi:.2f}] mm straddles the "
                f"{tl:.2f} mm requirement. This instrument will not issue a "
                f"verdict from an image in this case.")


@dataclass(frozen=True)
class OfficerDetermination:
    """The officer's call. Recorded as its own tier so it never looks machine-made."""

    officer_id: str
    verdict: str
    note: str = ""
    determined_on: Optional[datetime.date] = None


@dataclass(frozen=True)
class InspectionRecord:
    inspection_id: str
    inspected_on: datetime.date
    officer_id: str
    product_name: str
    findings: tuple[Finding, ...]
    coverage_description: str = ""
    extraction_provenance: tuple[str, ...] = ()
    measurement: Optional[Measurement] = None
    determination: Optional[OfficerDetermination] = None
    declared_category: str = ""
    declared_pdp_area_cm2: Optional[float] = None
    engine_versions: tuple[str, ...] = ()

    # ---- canonical form and hash ----------------------------------------
    def canonical_content(self) -> str:
        """A deterministic serialisation. The hash is taken over THIS.

        Deliberately excludes anything that varies between renders -- render
        time, file paths, PDF object ids. Two runs of the same inspection must
        produce the same digest or the digest means nothing.
        """
        parts: list[str] = [
            f"inspection_id={self.inspection_id}",
            f"inspected_on={self.inspected_on.isoformat()}",
            f"officer_id={self.officer_id}",
            f"product_name={self.product_name}",
            f"declared_category={self.declared_category}",
            f"declared_pdp_area_cm2={self.declared_pdp_area_cm2}",
            f"coverage={self.coverage_description}",
        ]
        for f_ in self.findings:
            parts.append("|".join([
                "finding", f_.key, f_.verdict, f_.source_tier,
                str(f_.detected), f_.why, f_.citation.printable(), f_.extractor]))
        m = self.measurement
        if m is not None:
            parts.append("|".join([
                "measurement", m.band, str(m.height_mm), str(m.u_mm),
                str(m.threshold_mm), m.convention, str(m.refusal_code),
                m.threshold_source_tier, str(m.roi_box)]))
            parts.extend(f"manifest|{line}" for line in m.manifest)
        d = self.determination
        if d is not None:
            parts.append("|".join([
                "determination", d.officer_id, d.verdict, d.note,
                d.determined_on.isoformat() if d.determined_on else "None"]))
        parts.extend(f"engine|{v}" for v in self.engine_versions)
        return "\n".join(parts)

    def content_hash(self) -> str:
        return hashlib.sha256(
            self.canonical_content().encode("utf-8")).hexdigest()

    def summary(self) -> dict:
        return summarise(self.findings)


HASH_NOTE = ("SHA-256 over the canonical inspection record, not over this "
             "file. Recompute it from an independently held copy of the record "
             "to detect alteration. This is a tamper-evident check, NOT a "
             "digital signature: it does not establish who produced this "
             "report.")

TIER_NOTE = {
    TIER_EXTRACTED: "Read from an image by OCR or a vision model. UNVERIFIED "
                    "-- confirm against the physical package before acting.",
    TIER_MEASURED: "Produced by a scale-referenced measurement with a stated "
                   "uncertainty (k=2). Traceability status is stated below.",
    TIER_DETERMINED: "The inspecting officer's own determination. Overrides "
                     "both of the above.",
}


# --------------------------------------------------------------------------
# Plain text
# --------------------------------------------------------------------------
def render_text(rec: InspectionRecord) -> str:
    s = rec.summary()
    L: list[str] = []
    add = L.append
    add("=" * 74)
    add("LEGAL METROLOGY (PACKAGED COMMODITIES) RULES, 2011")
    add("INSPECTION RECORD")
    add("=" * 74)
    add(f"Inspection ID : {rec.inspection_id}")
    add(f"Date          : {rec.inspected_on.isoformat()}")
    add(f"Officer       : {rec.officer_id}")
    add(f"Product       : {rec.product_name}")
    if rec.declared_category:
        add(f"Category      : {rec.declared_category}  (operator declared)")
    if rec.declared_pdp_area_cm2 is not None:
        add(f"PDP area      : {rec.declared_pdp_area_cm2} cm2  (operator declared)")
    add("")
    add(f"HEADLINE: {s['headline']}")
    add(f"  PASS {s['counts'][VERDICT_PASS]}    "
        f"FAIL {s['counts'][VERDICT_FAIL]}    "
        f"CANNOT DETERMINE {s['counts'][VERDICT_CANNOT_DETERMINE]}")
    add("")
    add("-" * 74)
    add("MANDATORY DECLARATIONS -- Rule 6(1)        [tier: EXTRACTED]")
    add("-" * 74)
    for f_ in rec.findings:
        add(f"[{f_.verdict}] {f_.label}")
        for line in f_.chain():
            add(f"      {line}")
        add("")

    if rec.measurement is not None:
        m = rec.measurement
        add("-" * 74)
        add("CHARACTER HEIGHT -- Rule 7(2)              [tier: MEASURED]")
        add("-" * 74)
        add(f"[{m.verdict()}] Height of the smallest governing character")
        add(f"      Measured:  {m.stated()}")
        if m.threshold_mm is not None:
            add(f"      Required:  {m.threshold_mm:.2f} mm")
        if m.threshold_source_tier:
            add(f"      Threshold source tier: {m.threshold_source_tier}")
        add(f"      Band:      {m.band}")
        add(f"      Why:       {m.band_explanation()}")
        if m.convention:
            add(f"      Convention: {m.convention}")
        if m.roi_box:
            add(f"      ROI:       {m.roi_box}  (operator-declared region)")
        if m.manifest:
            add("      Capture manifest:")
            for line in m.manifest:
                add(f"        {line}")
        add("")

    if rec.determination is not None:
        d = rec.determination
        add("-" * 74)
        add("OFFICER DETERMINATION                      [tier: DETERMINED]")
        add("-" * 74)
        add(f"[{d.verdict}] by {d.officer_id}"
            + (f" on {d.determined_on.isoformat()}" if d.determined_on else ""))
        if d.note:
            add(f"      {d.note}")
        add("")

    add("-" * 74)
    add("HOW EACH FINDING WAS OBTAINED")
    add("-" * 74)
    for tier in (TIER_EXTRACTED, TIER_MEASURED, TIER_DETERMINED):
        add(f"{tier}: {TIER_NOTE[tier]}")
    add("")
    if rec.coverage_description:
        add(f"Coverage: {rec.coverage_description}")
    for line in rec.extraction_provenance:
        # extraction_provenance_lines() leads with its own Coverage line; the
        # record already carries one, and printing both reads as a duplicated
        # claim rather than two sources agreeing.
        if line.strip().lower().startswith("coverage"):
            continue
        add(f"  {line}")
    add("")
    for v in rec.engine_versions:
        add(f"Engine: {v}")
    add("")
    add(f"Record digest (SHA-256): {rec.content_hash()}")
    add(HASH_NOTE)
    add("=" * 74)
    return "\n".join(L)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def render_pdf(rec: InspectionRecord, path: str) -> str:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm as MM
    from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    NAVY = colors.HexColor("#0A1628")
    GREY = colors.HexColor("#555555")
    RULE = colors.HexColor("#BBBBBB")

    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontName="Times-Bold",
                        fontSize=14, textColor=NAVY, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Times-Bold",
                        fontSize=10.5, textColor=NAVY, spaceBefore=10,
                        spaceAfter=4)
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Times-Roman",
                          fontSize=9, leading=12, alignment=TA_LEFT)
    small = ParagraphStyle("small", parent=body, fontSize=7.5, leading=9.5,
                           textColor=GREY)
    mono = ParagraphStyle("mono", parent=body, fontName="Courier", fontSize=7.5,
                          leading=9.5)

    doc = SimpleDocTemplate(path, pagesize=A4, title="Inspection Record",
                            leftMargin=18 * MM, rightMargin=18 * MM,
                            topMargin=16 * MM, bottomMargin=16 * MM)
    S: list = []
    s = rec.summary()

    S.append(Paragraph("Legal Metrology (Packaged Commodities) Rules, 2011", h1))
    S.append(Paragraph("INSPECTION RECORD", ParagraphStyle(
        "sub", parent=h1, fontSize=10, spaceAfter=8)))

    meta = [["Inspection ID", rec.inspection_id, "Date", rec.inspected_on.isoformat()],
            ["Officer", rec.officer_id, "Product", rec.product_name]]
    if rec.declared_category or rec.declared_pdp_area_cm2 is not None:
        meta.append(["Category (declared)", rec.declared_category or "-",
                     "PDP area (declared)",
                     f"{rec.declared_pdp_area_cm2} cm2"
                     if rec.declared_pdp_area_cm2 is not None else "-"])
    t = Table(meta, colWidths=[32 * MM, 52 * MM, 32 * MM, 58 * MM])
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Times-Roman"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("TEXTCOLOR", (0, 0), (0, -1), GREY),
        ("TEXTCOLOR", (2, 0), (2, -1), GREY),
        ("GRID", (0, 0), (-1, -1), 0.4, RULE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    S.append(t)
    S.append(Spacer(1, 8))

    S.append(Paragraph(f"<b>{s['headline']}</b> &nbsp;&nbsp; "
                       f"PASS {s['counts'][VERDICT_PASS]} &nbsp; "
                       f"FAIL {s['counts'][VERDICT_FAIL]} &nbsp; "
                       f"CANNOT DETERMINE {s['counts'][VERDICT_CANNOT_DETERMINE]}",
                       body))

    # ---- declarations ----------------------------------------------------
    S.append(Paragraph("Mandatory declarations &mdash; Rule 6(1) "
                       "<font size=7 color='#555555'>[tier: EXTRACTED]</font>", h2))
    rows = [["Verdict", "Declaration", "Detected", "Why", "Citation"]]
    for f_ in rec.findings:
        rows.append([Paragraph(f_.verdict.replace("_", "<br/>"), small),
                     Paragraph(f_.label, small),
                     Paragraph(str(f_.detected) if f_.detected else "&mdash;", small),
                     Paragraph(f_.why, small),
                     Paragraph(f_.citation.printable(), small)])
    dt = Table(rows, colWidths=[24 * MM, 34 * MM, 33 * MM, 60 * MM, 23 * MM],
               repeatRows=1)
    style = [("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
             ("FONTSIZE", (0, 0), (-1, -1), 7.5),
             ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8ECF1")),
             ("GRID", (0, 0), (-1, -1), 0.4, RULE),
             ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("TOPPADDING", (0, 0), (-1, -1), 3),
             ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
    for i, f_ in enumerate(rec.findings, start=1):
        if f_.verdict == VERDICT_FAIL:
            style.append(("BACKGROUND", (0, i), (0, i),
                          colors.HexColor("#F6E4E4")))
    dt.setStyle(TableStyle(style))
    S.append(dt)

    # ---- measurement: deliberately NOT the same table --------------------
    if rec.measurement is not None:
        m = rec.measurement
        S.append(Paragraph("Character height &mdash; Rule 7(2) "
                           "<font size=7 color='#555555'>[tier: MEASURED]</font>", h2))
        mrows = [["Verdict", m.verdict().replace("_", " ")],
                 ["Measured", m.stated()],
                 ["Required", f"{m.threshold_mm:.2f} mm"
                  if m.threshold_mm is not None else "not resolved"],
                 ["Threshold source", m.threshold_source_tier or "not stated"],
                 ["Basis", Paragraph(m.band_explanation(), small)]]
        if m.convention:
            mrows.append(["Convention", Paragraph(m.convention, small)])
        if m.roi_box:
            mrows.append(["Region measured",
                          f"{m.roi_box}  (operator-declared)"])
        mt = Table(mrows, colWidths=[34 * MM, 140 * MM])
        mt.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (0, -1), "Times-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3F0E7")),
            ("BOX", (0, 0), (-1, -1), 1.1, NAVY),
            ("INNERGRID", (0, 0), (-1, -1), 0.4, RULE),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        S.append(mt)
        if m.manifest:
            S.append(Spacer(1, 4))
            S.append(Paragraph("Capture manifest", ParagraphStyle(
                "cm", parent=small, fontName="Times-Bold")))
            for line in m.manifest:
                S.append(Paragraph(line.replace("&", "&amp;"), mono))

    # ---- determination ---------------------------------------------------
    if rec.determination is not None:
        d = rec.determination
        S.append(Paragraph("Officer determination "
                           "<font size=7 color='#555555'>[tier: DETERMINED]</font>", h2))
        S.append(Paragraph(
            f"<b>{d.verdict}</b> &mdash; {d.officer_id}"
            + (f", {d.determined_on.isoformat()}" if d.determined_on else "")
            + (f"<br/>{d.note}" if d.note else ""), body))

    # ---- provenance ------------------------------------------------------
    S.append(Paragraph("How each finding was obtained", h2))
    for tier in (TIER_EXTRACTED, TIER_MEASURED, TIER_DETERMINED):
        S.append(Paragraph(f"<b>{tier}</b> &mdash; {TIER_NOTE[tier]}", small))
    if rec.coverage_description:
        S.append(Spacer(1, 3))
        S.append(Paragraph(f"<b>Coverage</b> &mdash; {rec.coverage_description}", small))
    for line in rec.extraction_provenance:
        if line.strip().lower().startswith("coverage"):
            continue      # already printed above; see _dedupe note
        S.append(Paragraph(line.replace("&", "&amp;"), small))

    S.append(Spacer(1, 6))
    for v in rec.engine_versions:
        S.append(Paragraph(f"Engine: {v}", small))
    S.append(Spacer(1, 4))
    S.append(KeepTogether([
        Paragraph(f"<b>Record digest (SHA-256)</b>", small),
        Paragraph(rec.content_hash(), mono),
        Paragraph(HASH_NOTE, small)]))

    doc.build(S)
    return path


# --------------------------------------------------------------------------
# DOCX -- the problem statement's "editable format"
# --------------------------------------------------------------------------
def render_docx(rec: InspectionRecord, path: str) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    NAVY = RGBColor(0x0A, 0x16, 0x28)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Legal Metrology (Packaged Commodities) Rules, 2011")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY
    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = t2.add_run("INSPECTION RECORD")
    r2.bold = True
    r2.font.color.rgb = NAVY

    meta = doc.add_table(rows=0, cols=4)
    meta.style = "Table Grid"
    for a, b, c, d in (("Inspection ID", rec.inspection_id, "Date",
                        rec.inspected_on.isoformat()),
                       ("Officer", rec.officer_id, "Product", rec.product_name),
                       ("Category (declared)", rec.declared_category or "-",
                        "PDP area (declared)",
                        f"{rec.declared_pdp_area_cm2} cm2"
                        if rec.declared_pdp_area_cm2 is not None else "-")):
        cells = meta.add_row().cells
        for i, v in enumerate((a, b, c, d)):
            cells[i].text = str(v)

    s = rec.summary()
    p = doc.add_paragraph()
    p.add_run(s["headline"]).bold = True
    p.add_run(f"    PASS {s['counts'][VERDICT_PASS]}   "
              f"FAIL {s['counts'][VERDICT_FAIL]}   "
              f"CANNOT DETERMINE {s['counts'][VERDICT_CANNOT_DETERMINE]}")

    doc.add_heading("Mandatory declarations — Rule 6(1)  [tier: EXTRACTED]", level=2)
    dt = doc.add_table(rows=1, cols=5)
    dt.style = "Table Grid"
    for i, h in enumerate(("Verdict", "Declaration", "Detected", "Why", "Citation")):
        cell = dt.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
    for f_ in rec.findings:
        c = dt.add_row().cells
        c[0].text = f_.verdict.replace("_", " ")
        c[1].text = f_.label
        c[2].text = str(f_.detected) if f_.detected else "—"
        c[3].text = f_.why
        c[4].text = f_.citation.printable()

    if rec.measurement is not None:
        m = rec.measurement
        doc.add_heading("Character height — Rule 7(2)  [tier: MEASURED]", level=2)
        mt = doc.add_table(rows=0, cols=2)
        mt.style = "Table Grid"
        rows = [("Verdict", m.verdict().replace("_", " ")),
                ("Measured", m.stated()),
                ("Required", f"{m.threshold_mm:.2f} mm"
                 if m.threshold_mm is not None else "not resolved"),
                ("Threshold source", m.threshold_source_tier or "not stated"),
                ("Basis", m.band_explanation())]
        if m.convention:
            rows.append(("Convention", m.convention))
        if m.roi_box:
            rows.append(("Region measured", f"{m.roi_box}  (operator-declared)"))
        for k, v in rows:
            c = mt.add_row().cells
            c[0].text = ""
            c[0].paragraphs[0].add_run(k).bold = True
            c[1].text = str(v)
        if m.manifest:
            doc.add_paragraph().add_run("Capture manifest").bold = True
            for line in m.manifest:
                doc.add_paragraph(line, style="List Bullet")

    if rec.determination is not None:
        d = rec.determination
        doc.add_heading("Officer determination  [tier: DETERMINED]", level=2)
        p = doc.add_paragraph()
        p.add_run(d.verdict).bold = True
        p.add_run(f" — {d.officer_id}"
                  + (f", {d.determined_on.isoformat()}" if d.determined_on else ""))
        if d.note:
            doc.add_paragraph(d.note)

    doc.add_heading("How each finding was obtained", level=2)
    for tier in (TIER_EXTRACTED, TIER_MEASURED, TIER_DETERMINED):
        p = doc.add_paragraph()
        p.add_run(f"{tier} — ").bold = True
        p.add_run(TIER_NOTE[tier])
    if rec.coverage_description:
        p = doc.add_paragraph()
        p.add_run("Coverage — ").bold = True
        p.add_run(rec.coverage_description)
    for line in rec.extraction_provenance:
        if line.strip().lower().startswith("coverage"):
            continue
        doc.add_paragraph(line)
    for v in rec.engine_versions:
        doc.add_paragraph(f"Engine: {v}")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Record digest (SHA-256)").bold = True
    doc.add_paragraph(rec.content_hash())
    doc.add_paragraph(HASH_NOTE)

    doc.save(path)
    return path


# --------------------------------------------------------------------------
# Self-test
# --------------------------------------------------------------------------
def _sample_record(with_measurement: bool = True,
                   refused: bool = False) -> InspectionRecord:
    from lm_declarations import (DECL_COMMON_NAME, DECL_CONSUMER_CARE,
                                 DECL_DATE, DECL_NAME_ADDRESS,
                                 DECL_NET_QUANTITY, DECL_RETAIL_PRICE,
                                 check_declarations)
    from lm_extract import Coverage, extraction_provenance_lines, manual_extract

    cov = Coverage(panels=("front", "back"), operator_examined_package=True,
                   operator_id="INS-01", note="copy stand, controlled light")
    fields = manual_extract({
        DECL_NAME_ADDRESS: "Acme Foods Pvt Ltd, Chennai 600001",
        DECL_COMMON_NAME: "Biscuits",
        DECL_NET_QUANTITY: "100 g",
        DECL_DATE: "08/2025",
        DECL_RETAIL_PRICE: "Rs. 250",
        DECL_CONSUMER_CARE: None,
    }, cov, "INS-01")

    if refused:
        m = Measurement(band=BAND_REFER, refusal_code="THRESHOLD_DISPUTED",
                        refusal_detail="The applicable height requirement for "
                                       "this panel area is not settled between "
                                       "sources. This instrument will not rule.",
                        threshold_source_tier="AGGREGATOR / DISPUTED",
                        convention="50% ink-to-substrate crossing")
    else:
        m = Measurement(band=BAND_REFER, height_mm=1.1538, u_mm=0.1690,
                        threshold_mm=1.0,
                        convention="50% ink-to-substrate crossing, "
                                   "max-supported extent",
                        threshold_source_tier="PRIMARY_GAZETTE / VERIFIED",
                        roi_box=(412, 233, 690, 96),
                        manifest=("scale: 121.4 px/mm from chessboard target",
                                  "artifact tier: PRINTED_SPECIMEN (not a standard)",
                                  "focus: 212.4 (floor 180.0)"))

    return InspectionRecord(
        inspection_id="VEC-2026-0912-001",
        inspected_on=datetime.date(2026, 9, 12),
        officer_id="INS-01",
        product_name="Biscuits 100 g",
        findings=check_declarations(fields),
        coverage_description=cov.describe(),
        extraction_provenance=extraction_provenance_lines(fields, cov),
        measurement=m if with_measurement else None,
        determination=OfficerDetermination(
            officer_id="INS-01", verdict="REFERRED FOR PHYSICAL VERIFICATION",
            note="Consumer care absent on examination. Height indeterminate "
                 "from image; packet retained for physical verification.",
            determined_on=datetime.date(2026, 9, 12)),
        declared_category="general",
        declared_pdp_area_cm2=40.0,
        engine_versions=("lm_declarations 1.0-chunk1-2026-09-12",
                         "lm_metrology 7.9-call-site-3-2026-09-07"))


def _self_test() -> int:
    checks = 0
    fails: list[str] = []

    def ck(name: str, cond: bool, detail: str = "") -> None:
        nonlocal checks
        checks += 1
        if not cond:
            fails.append(name + (f" -- {detail}" if detail else ""))

    rec = _sample_record()

    # --- hash is over content, and is stable -------------------------------
    ck("hash is stable across calls", rec.content_hash() == rec.content_hash())
    ck("hash is stable across identical records",
       _sample_record().content_hash() == rec.content_hash())
    other = _sample_record()
    changed = InspectionRecord(**{**other.__dict__, "product_name": "Something else"})
    ck("hash changes when content changes",
       changed.content_hash() != rec.content_hash())
    ck("hash is a sha256 hex digest", len(rec.content_hash()) == 64)
    ck("hash note does not claim to be a signature",
       "NOT a digital signature" in HASH_NOTE)

    # --- the measurement invariant ----------------------------------------
    try:
        Measurement(band=BAND_COMPLIANT, height_mm=1.2, u_mm=None)
        ck("height without uncertainty must raise", False, "did not raise")
    except ValueError:
        ck("height without uncertainty raises", True)
    try:
        Measurement(band=BAND_COMPLIANT, height_mm=None, u_mm=0.1)
        ck("uncertainty without height must raise", False, "did not raise")
    except ValueError:
        ck("uncertainty without height raises", True)

    # --- no score anywhere -------------------------------------------------
    # A bare "%" is not the thing to forbid: the measurement convention is
    # legitimately "50% ink-to-substrate crossing". What must never appear is a
    # percentage or grade attached to COMPLIANCE, because that would imply the
    # six declarations are commensurable and that CANNOT_DETERMINE can be
    # averaged with PASS.
    txt = render_text(rec)
    import re as _re
    ck("no compliance percentage",
       not _re.search(r"\d+\s*%\s*(compliant|compliance|pass|score)", txt, _re.I))
    ck("no 'compliance score' phrasing",
       not _re.search(r"(compliance|overall)\s+(score|grade|rating)", txt, _re.I))
    ck("the legitimate 50% convention survives", "50% ink-to-substrate" in txt,
       "the test must not be so blunt it forbids the measurand definition")
    for word in ("accuracy",):
        ck(f"the word {word!r} does not appear", word not in txt.lower(),
           "this system has not measured a reference standard")

    # --- the disputed bracket must not leak candidate values ---------------
    ref = _sample_record(refused=True)
    rtxt = render_text(ref)
    for leaked in ("1.5", "2.0"):
        ck(f"disputed refusal does not leak {leaked!r}", leaked not in rtxt)
    ck("a refusal prints no height", "not measured" in rtxt)
    ck("a refusal names its code", "THRESHOLD_DISPUTED" in rtxt)

    # --- three tiers are all present and distinguishable -------------------
    for tier in (TIER_EXTRACTED, TIER_MEASURED, TIER_DETERMINED):
        ck(f"tier {tier} appears in the report", tier in txt)
    ck("the measured row states an uncertainty", "(k=2)" in txt)
    ck("the extracted rows are marked unverified",
       "UNVERIFIED" in txt.upper())
    ck("the threshold's source tier is printed",
       "PRIMARY_GAZETTE" in txt or "AGGREGATOR" in txt)

    # --- record without a measurement still renders ------------------------
    nom = _sample_record(with_measurement=False)
    ntxt = render_text(nom)
    ck("a record with no measurement renders", len(ntxt) > 200)
    ck("and does not fabricate a height section",
       "[tier: MEASURED]" not in ntxt)

    # --- the coverage claim appears once, not twice ------------------------
    # Found 2026-09-12 by rendering the PDF and looking at it: the record
    # carries a coverage description AND extraction_provenance_lines() leads
    # with its own. Both were printed. Two identical claims side by side read
    # as a duplicated assertion rather than as corroboration.
    ck("coverage is stated exactly once",
       txt.lower().count("physical package examined by") == 1,
       f"appears {txt.lower().count('physical package examined by')} times")

    # --- the two decimal places rule --------------------------------------
    m = rec.measurement
    ck("height is stated to 2 d.p.", m.stated().startswith("1.15 mm"), m.stated())
    ck("full precision is NOT printed", "1.1538" not in txt)

    print(f"report self-test: {checks} checks, {len(fails)} failed")
    for f_ in fails:
        print("  !", f_)
    return 1 if fails else 0


if __name__ == "__main__":
    import os
    print(f"lm_report  VERSION {VERSION}")
    print()
    rc = _self_test()
    print()
    rec = _sample_record()
    out = os.environ.get("LM_REPORT_OUT", ".")
    p = render_pdf(rec, os.path.join(out, "inspection-record-sample.pdf"))
    d = render_docx(rec, os.path.join(out, "inspection-record-sample.docx"))
    print(f"wrote {p}")
    print(f"wrote {d}")
    print()
    print(render_text(rec))
    raise SystemExit(rc)
